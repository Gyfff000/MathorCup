"""AI-assisted implementation: OpenAI Codex / GPT-6, 2026-10-24."""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .common import progress, save_csv, sha256, write_json
from .data import META


def run_audit(data, out, cfg):
    dest = Path(out) / "audit"
    dest.mkdir(parents=True, exist_ok=True)
    progress("Audit: checking source hashes, sample keys, scales and sensory templates")
    hashes, shapes = [], []
    for p in sorted(data.root.iterdir()):
        if p.is_file():
            hashes.append({"file": p.name, "bytes": p.stat().st_size, "sha256": sha256(p)})
            if p.suffix == ".csv":
                frame = pd.read_csv(p)
                shapes.append({"file": p.name, "rows": len(frame), "columns_including_metadata": len(frame.columns), "missing_cells": int(frame.isna().sum().sum())})
    save_csv(pd.DataFrame(hashes), dest / "source_manifest.csv")
    save_csv(pd.DataFrame(shapes), dest / "file_shapes.csv")
    save_csv(data.meta.groupby("tasting_category_fine").size().rename("n").reset_index(), dest / "style_counts.csv")
    save_csv(data.chemical_meta, dest / "chemical_dictionary.csv")
    names = data.sensory_names[data.sensory_names.dataset == "Trained panel"].drop_duplicates("R_name").set_index("R_name")
    targets = []
    for c in data.Y:
        alias = "X4vg" if c == "X4vg" else c
        targets.append({"field": c, "label": names.at[alias, "CleanName_PanelFeedback"] if alias in names.index else c,
                        "group": "Aroma" if c.startswith("A_") else "Flavour" if c.startswith("F_") else "Other sensory",
                        "observed_scale": "provided centered/standardized continuous values",
                        "absolute_0_6_calibration": "unresolved"})
    save_csv(pd.DataFrame(targets), dest / "sensory_dictionary.csv")
    records = []
    for col in data.X:
        a = data.X[col]
        raw = data.calibrated[col] if col in data.calibrated else None
        row = {"field": col, "n_missing": int(a.isna().sum()), "minimum": float(a.min()), "maximum": float(a.max()),
               "q01": float(a.quantile(.01)), "q99": float(a.quantile(.99)), "transform_evidence": "not_available",
               "raw_identity_fraction": np.nan, "raw_log10_fraction": np.nan, "n_raw_positive_pairs": 0}
        if raw is not None:
            mask = a.notna() & raw.notna() & (raw > 0)
            if mask.sum():
                identity = np.isclose(a[mask], raw[mask], atol=1e-7, rtol=1e-6)
                logarithm = np.isclose(a[mask], np.log10(raw[mask]), atol=1e-7, rtol=1e-6)
                fi, fl = float(identity.mean()), float(logarithm.mean())
                row.update(raw_identity_fraction=fi, raw_log10_fraction=fl, n_raw_positive_pairs=int(mask.sum()))
                row["transform_evidence"] = "verified_log10_positive_pairs" if fl >= .999 else "verified_identity_positive_pairs" if fi >= .999 else "mostly_identity_with_mismatches" if fi >= .9 else "mixed_or_unresolved"
        records.append(row)
    scales = pd.DataFrame(records)
    save_csv(scales, dest / "field_scales.csv")
    # Inspect actual input cells, not formula caches that evaluate blanks to zero.
    w = load_workbook(data.root / "3.xlsx", data_only=False)
    input_rows = [7, 20, 33, 45, 50, 54, 58, 62, 68, 72, 76, 82, 86, 90, 94]
    input_cols = [3, 5, 7, 9, 11, 13, 15, 19, 21, 23, 25, 27, 29, 31]
    template = []
    for sh in w:
        if not sh.title.startswith("beer_"):
            continue
        marks = [sh.cell(r, c).coordinate for r in input_rows for c in input_cols if sh.cell(r, c).value is not None]
        template.append({"sheet": sh.title, "nonempty_main_scoring_inputs": len(marks), "input_cells": ",".join(marks)})
    save_csv(pd.DataFrame(template), dest / "sensory_template_audit.csv")
    key_stats = []
    for n in [1, 4, 7]:
        d = pd.read_csv(data.root / f"{n}.csv")
        key_stats.append({"file": f"{n}.csv", "unique_beer": int(d.beer.nunique()), "unique_beer_id": int(d.beer_id.nunique())})
    save_csv(pd.DataFrame(key_stats), dest / "join_key_audit.csv")
    values = data.given_corr.to_numpy()
    evidence = {
        "pearson_matrix_mean_absolute_difference": float(np.nanmean(np.abs(data.X.corr().to_numpy() - values))),
        "spearman_matrix_mean_absolute_difference": float(np.nanmean(np.abs(data.X.corr(method="spearman").to_numpy() - values))),
        "n_samples": len(data.X), "n_features": data.X.shape[1], "n_targets": data.Y.shape[1],
        "n_styles": int(data.meta.tasting_category_fine.nunique()),
        "grade_calibration": "UNRESOLVED: no observed paired absolute 0-6 anchors confirmed",
        "template_summary_barnyard_formula": str(w["summary"]["AX2"].value),
        "template_barnyard_expected_cell": "beer_1!AG90 based on the form's Barnyard label; team review required",
        "scope": "source evidence verified by computation; model conclusions remain exploratory",
    }
    write_json(dest / "verified_facts.json", evidence)
