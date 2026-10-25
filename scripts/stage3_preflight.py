"""Read-only source audit and conditional analysis of existing OOF predictions.

Run from the bundled document runtime (openpyxl + numpy). No workbook is edited.
This is a protocol preflight, not the full stage-three experiment runner.
"""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.stage3_algorithms import paired_oof_bootstrap

parser = argparse.ArgumentParser()
parser.add_argument("--data-dir", type=Path, default=ROOT.parent / "materials/A/附件")
parser.add_argument("--out", type=Path, default=ROOT / "outputs/stage3_design")
args = parser.parse_args()
DATA, OUT = args.data_dir, args.out
OUT.mkdir(parents=True, exist_ok=True)


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def write(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


original_hashes = {name: hashlib.sha256((DATA / name).read_bytes()).hexdigest() for name in ["1.csv", "3.xlsx", "4.csv", "8.xlsx"]}
previous_hashes = {r['file']:r['sha256'] for r in read_csv(ROOT/'outputs/stage2/audit/source_manifest.csv')}
if any(previous_hashes.get(name) != sha for name,sha in original_hashes.items()):
    raise ValueError("Source version differs from the existing stage-two predictions/evidence")
target_rows = read_csv(DATA / "4.csv")
targets = [c for c in target_rows[0] if c not in ["beer", "beer_id", "tasting_category_fine"]]
aliases = {"A_ester_fruity": "A_esters_fruity", "F_ester_fruity": "F_esters_fruity", "4vg": "X4vg"}
workbook = load_workbook(DATA / "3.xlsx", data_only=False, read_only=True)
summary = workbook["summary"]
mapping, anomalies = [], []
for col in range(2, summary.max_column+1):
    original = summary.cell(1,col).value
    target = aliases.get(original,original)
    if target not in targets:
        continue
    formula = str(summary.cell(2,col).value)
    match = re.fullmatch(r"=beer_1!\$([A-Z]+)\$(\d+)",formula)
    if not match:
        raise ValueError(f"Unexpected score formula for {target}: {formula}")
    column, row = match.group(1), int(match.group(2))
    expected_cell = "AG90" if target == "barnyard" else f"{column}{row}"
    section = ("Aroma_core" if target.startswith("A_") else "Flavour_core" if target.startswith("F_") else
               "Faults" if target in ["dms","metallic","X4vg","t2n","diacetyl","stale_hops"] else
               "Spiced_beers" if target in ["orange","coriander","clove"] else
               "Lambic_Geuze" if target in ["lactic","acetic","barnyard"] else
               "Overall" if target == "overall" else "Palate_or_basic_taste")
    mapping.append({"target":target,"summary_header":original,"summary_column":col,"original_formula":formula,
                    "expected_score_cell":expected_cell,"section":section,
                    "formula_status":"suspected_wrong_reference_derived_map_corrects_only" if target=='barnyard' else "consistent_with_form_position",
                    "official_submission_scope":"requires_explicit_reporting_core_30_and_full_50_not_inferred_from_count"})
    if target == "barnyard":
        assert "Barnyard" in str(workbook["beer_1"]["R89"].value)
        assert str(workbook["beer_1"]["AG90"].value).startswith("=")
        for i in range(2,14):
            actual=str(summary.cell(i,col).value)
            expected=f"=beer_{i-1}!$AG$90"
            if actual != expected:
                anomalies.append({"summary_cell":summary.cell(i,col).coordinate,"target":target,"actual_formula":actual,"form_position_expected_formula":expected})
assert len(mapping) == len(targets) == 50 and {r['target'] for r in mapping} == set(targets)
write("scoring_field_map.json",mapping)
write("scoring_scope_audit.json",{"mapped_targets":50,"core_Aroma":15,"core_Flavour":15,"additional_targets":20,
    "aliases":aliases,"suspected_formula_mismatches":anomalies,
    "card_items_without_separate_supervised_targets":[{"card_label":"Acetaldehyde under Aroma","score_cell":"Q45"},{"card_label":"Acetaldehyde under Flavour","score_cell":"AG45"}],
    "original_files_edited":False,"absolute_score_anchors_available":False,
    "scope_caveat":"30 is the prefixed core comparison scope, not a ruling that only 30 answers are required."})

scales = {r['field']:r for r in read_csv(ROOT / "outputs/stage2/audit/field_scales.csv")}
chem = {r['field']:r for r in read_csv(ROOT / "outputs/stage2/audit/chemical_dictionary.csv")}
registry=[]
for field in ["DMS","X2.3.butanedione","X4VG","X.E..2.nonenal"]:
    scale,meta=scales[field],chem[field]
    try: threshold=float(meta['threshold_median'])
    except (TypeError,ValueError): threshold=None
    if threshold is not None and not np.isfinite(threshold): threshold=None
    supported=scale['transform_evidence']=='verified_log10_positive_pairs'
    registry.append({"field":field,"unit_in_attachment":meta['unit'],"threshold_median_in_attachment":threshold,
        "transformation_evidence":scale['transform_evidence'],"positive_pair_log10_supported":supported,
        "concentration_rule_if_eligible":"10**x" if supported else "unresolved_per_sample_check_needed",
        "zero_LOD_or_substitution_provenance":"not_fully_verified",
        "absolute_defect_certification":"not_complete_even_if_positive_pairs_match",
        "next_action":"audit_per_sample_calibrated_pairs_and_detection_limits" if supported else "resolve_mismatching_pairs_and_scale_before_concentration_claim"})
write("defect_registry_preflight.json",registry)

beer_ids=[int(r['beer']) for r in target_rows]
truth=np.array([[float(r[t]) for t in targets] for r in target_rows])
def predictions(path):
    rows=read_csv(path)
    bybeer={int(r['beer']):r for r in rows}
    if len(bybeer)!=len(rows) or set(bybeer)!=set(beer_ids): raise ValueError("OOF keys do not align uniquely")
    return np.array([[float(bybeer[b][t]) for t in targets] for b in beer_ids])
gbr=predictions(ROOT/'outputs/stage2/q1/oof_gbr_continuous.csv')
et=predictions(ROOT/'outputs/stage1/q1/oof_extra_trees_continuous.csv')
labels=[r['tasting_category_fine'] for r in target_rows]
reports=[]
for scope,indices in [('all_50',list(range(50))),('prefixed_core_30',[j for j,t in enumerate(targets) if t.startswith(('A_','F_'))])]:
    result=paired_oof_bootstrap(truth[:,indices],gbr[:,indices],et[:,indices],labels)
    result.update(scope=scope,comparison='fixed_stage2_GBR_minus_stage1_ExtraTrees',metric='macro_MAE',
                  uncertainty_scope='beer-resampling_conditional_on_existing_predictions_not_fresh_external_validation')
    reports.append(result)
write('existing_oof_paired_pilot.json',reports)

test_report=ROOT/'tmp/stage3_tests.xml'
suite=ET.parse(test_report).getroot().find('testsuite')
assert all(int(suite.attrib[k])==0 for k in ['failures','errors','skipped'])
for name,old in original_hashes.items():
    assert hashlib.sha256((DATA/name).read_bytes()).hexdigest()==old
manifest={'date':'2026-10-25','stage':'design_and_preflight','complete_stage3_training_run':False,
    'source_sha256':original_hashes,'protocol_sha256':hashlib.sha256((ROOT/'config_stage3_design.json').read_bytes()).hexdigest(),
    'algorithm_tests':{k:int(suite.attrib[k]) for k in ['tests','failures','errors','skipped']},
    'no_human_validation_claimed':True,'no_original_workbook_edited':True,
    'full_experiments_status':'planned_not_executed; only source mapping and conditional existing-OOF pilot executed'}
write('preflight_manifest.json',manifest)
print(json.dumps({'mapped_targets':len(mapping),'formula_mismatches':len(anomalies),'paired_pilot':reports,'test_summary':manifest['algorithm_tests']},ensure_ascii=False))
